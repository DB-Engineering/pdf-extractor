#!/usr/bin/env python3
"""
PDF & Table Extractor (PDF to Text & Excel Converter)
=====================================================
Extracts text from PDF files and converts embedded tables/schedules
into formatted multi-tab Excel spreadsheets (.xlsx).

Directory Structure:
--------------------
root/
├── pdf_to_excel.py
├── pdf_files/        (Input PDFs, e.g. page_1.pdf)
├── text_files/       (Output/Input text files, e.g. page_1.txt)
└── excel_files/      (Output Excel files, e.g. page_1.xlsx)

Usage:
    # Process a specific PDF (looks in current dir or pdf_files/):
    python3 pdf_to_excel.py page_1.pdf

    # Process all PDFs in the pdf_files/ directory:
    python3 pdf_to_excel.py

    # Only extract text to text_files/:
    python3 pdf_to_excel.py page_1.pdf --text-only

    # Only extract tables to excel_files/:
    python3 pdf_to_excel.py page_1.pdf --excel-only
"""

import argparse
import re
import sys
import time
from pathlib import Path

# Try importing PDF extraction libraries
try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    import pdfplumber
    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False

# Try importing openpyxl for Excel output
try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# Default headers for terminal schedules
DEFAULT_VAV_HEADERS = [
    "UNIT MARK", "MANUFACTURER", "MODEL", "INLET SIZE",
    "CFM MAX", "CFM MIN", "HEAT COIL MBH", "ROWS",
    "GPM", "3 WAY VALVE", "CO2 SENSOR", "OCC SENSOR", "REMARKS"
]


def setup_directories(base_dir: Path) -> tuple[Path, Path, Path]:
    """Ensures pdf_files, text_files, and excel_files directories exist."""
    pdf_dir = base_dir / "pdf_files"
    text_dir = base_dir / "text_files"
    excel_dir = base_dir / "excel_files"

    pdf_dir.mkdir(parents=True, exist_ok=True)
    text_dir.mkdir(parents=True, exist_ok=True)
    excel_dir.mkdir(parents=True, exist_ok=True)

    return pdf_dir, text_dir, excel_dir


def resolve_input_file(file_name: str, base_dir: Path) -> Path:
    """
    Resolves an input file path.
    Checks relative to working dir, then in pdf_files/ or text_files/.
    """
    pdf_dir, text_dir, _ = setup_directories(base_dir)
    target = Path(file_name)

    # 1. Check direct path
    if target.exists() and target.is_file():
        return target.resolve()

    # 2. Check in pdf_files/
    in_pdf = pdf_dir / target.name
    if in_pdf.exists() and in_pdf.is_file():
        return in_pdf.resolve()

    # 3. Check in text_files/
    in_text = text_dir / target.name
    if in_text.exists() and in_text.is_file():
        return in_text.resolve()

    # 4. Try adding .pdf extension
    if not target.suffix:
        in_pdf_ext = pdf_dir / f"{target.name}.pdf"
        if in_pdf_ext.exists():
            return in_pdf_ext.resolve()

        in_text_ext = text_dir / f"{target.name}.txt"
        if in_text_ext.exists():
            return in_text_ext.resolve()

    raise FileNotFoundError(
        f"Could not locate file '{file_name}'. Checked:\n"
        f"  • Direct path: {target.resolve()}\n"
        f"  • PDF directory: {pdf_dir / target.name}\n"
        f"  • Text directory: {text_dir / target.name}"
    )


def extract_text_pypdf(pdf_path: Path, include_headers: bool = True) -> tuple[str, int]:
    """Extract text from PDF using pypdf."""
    reader = pypdf.PdfReader(str(pdf_path))
    total_pages = len(reader.pages)
    text_blocks = []

    for idx, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if include_headers:
            header = f"--- Page {idx} of {total_pages} ---"
            text_blocks.append(f"{header}\n\n{page_text.strip()}\n")
        else:
            text_blocks.append(page_text.strip())

    full_text = "\n\n".join(text_blocks)
    return full_text, total_pages


def extract_text_pdfplumber(pdf_path: Path, include_headers: bool = True) -> tuple[str, int]:
    """Extract text from PDF using pdfplumber."""
    text_blocks = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        total_pages = len(pdf.pages)
        for idx, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            if include_headers:
                header = f"--- Page {idx} of {total_pages} ---"
                text_blocks.append(f"{header}\n\n{page_text.strip()}\n")
            else:
                text_blocks.append(page_text.strip())

    full_text = "\n\n".join(text_blocks)
    return full_text, total_pages


def convert_pdf_to_text(
    pdf_path: str | Path,
    output_path: str | Path | None = None,
    base_dir: Path | None = None,
    engine: str = "auto",
    no_page_headers: bool = False,
    quiet: bool = False
) -> Path:
    """Converts a PDF file to a plain text (.txt) file in text_files/."""
    script_dir = base_dir or Path(__file__).resolve().parent
    _, text_dir, _ = setup_directories(script_dir)

    input_file = Path(pdf_path).resolve()
    if not input_file.exists():
        raise FileNotFoundError(f"Input PDF file not found at '{input_file}'")

    if output_path is None:
        output_file = text_dir / f"{input_file.stem}.txt"
    else:
        output_file = Path(output_path).resolve()

    selected_engine = engine.lower()
    if selected_engine == "auto":
        if PYPDF_AVAILABLE:
            selected_engine = "pypdf"
        elif PDFPLUMBER_AVAILABLE:
            selected_engine = "pdfplumber"
        else:
            raise RuntimeError(
                "No supported PDF library installed. Please install 'pypdf':\n"
                "  • Via pip: pip install pypdf --break-system-packages\n"
                "  • Via apt: sudo apt install python3-pypdf"
            )

    start_time = time.time()
    include_headers = not no_page_headers

    if selected_engine == "pypdf":
        extracted_text, total_pages = extract_text_pypdf(input_file, include_headers=include_headers)
    elif selected_engine == "pdfplumber":
        extracted_text, total_pages = extract_text_pdfplumber(input_file, include_headers=include_headers)
    else:
        raise ValueError(f"Unsupported engine '{selected_engine}'. Use 'pypdf' or 'pdfplumber'.")

    output_file.parent.mkdir(parents=True, exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as f:
        f.write(extracted_text)

    elapsed = time.time() - start_time
    if not quiet:
        rel_input = input_file.relative_to(script_dir) if input_file.is_relative_to(script_dir) else input_file
        rel_output = output_file.relative_to(script_dir) if output_file.is_relative_to(script_dir) else output_file
        print(f"Text Extraction Successful!")
        print(f"  • Input PDF:    {rel_input}")
        print(f"  • Output Text:  {rel_output}")
        print(f"  • Total Pages:  {total_pages}")
        print(f"  • Elapsed Time: {elapsed:.2f}s")

    return output_file


def parse_tables_from_text(text_content: str) -> list[dict]:
    """
    Parses extracted text into structured table records.
    Groups data rows under detected schedule titles (e.g., 1ST FLOOR, 3RD FLOOR).
    """
    lines = [line.strip() for line in text_content.splitlines() if line.strip()]
    tables = []

    schedule_titles = []
    for line in lines:
        if any(keyword in line.upper() for keyword in ["SINGLE DUCT", "TERMINAL SCEHDULE", "TERMINAL SCHEDULE", "EQUIPMENT SCHEDULE"]):
            schedule_titles.append(line.strip())

    vav_blocks = []
    temp_rows = []

    for line in lines:
        if line.startswith("VAV-") or re.match(r'^VAV\d*', line):
            tokens = line.split()
            temp_rows.append(tokens)
        elif any(keyword in line.upper() for keyword in ["SINGLE DUCT", "TERMINAL SCEHDULE", "TERMINAL SCHEDULE", "EQUIPMENT SCHEDULE"]):
            title = line.strip()
            if temp_rows:
                vav_blocks.append((title, temp_rows))
                temp_rows = []

    if temp_rows and not vav_blocks:
        title = schedule_titles[0] if schedule_titles else "Terminal Schedule"
        vav_blocks.append((title, temp_rows))

    for title, rows in vav_blocks:
        tables.append({
            "title": title,
            "headers": DEFAULT_VAV_HEADERS,
            "rows": rows
        })

    return tables


def extract_tables_from_pdf_or_text(file_path: Path) -> list[dict]:
    """
    Extracts tables from a PDF using pdfplumber's native table finder,
    falling back to text-based schedule parsing if needed.
    """
    tables = []

    if file_path.suffix.lower() == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text_content = f.read()
        return parse_tables_from_text(text_content)

    if PDFPLUMBER_AVAILABLE:
        try:
            with pdfplumber.open(str(file_path)) as pdf:
                for page_idx, page in enumerate(pdf.pages, start=1):
                    extracted = page.extract_tables()
                    for t_idx, raw_table in enumerate(extracted, start=1):
                        if not raw_table or len(raw_table) < 2:
                            continue
                        headers = [str(cell or "").strip() for cell in raw_table[0]]
                        rows = [[str(cell or "").strip() for cell in row] for row in raw_table[1:]]
                        tables.append({
                            "title": f"Page {page_idx} Table {t_idx}",
                            "headers": headers,
                            "rows": rows
                        })
        except Exception:
            pass

    if not tables:
        if PDFPLUMBER_AVAILABLE:
            text_content, _ = extract_text_pdfplumber(file_path, include_headers=False)
        elif PYPDF_AVAILABLE:
            text_content, _ = extract_text_pypdf(file_path, include_headers=False)
        else:
            raise RuntimeError("No PDF library available for extraction.")

        tables = parse_tables_from_text(text_content)

    return tables


def sanitize_sheet_title(raw_title: str, index: int) -> str:
    """Formats a raw schedule title into a valid Excel sheet name (max 31 chars)."""
    clean = re.sub(r'[\\/*?:\[\]]', '', raw_title).strip()

    floor_match = re.search(r'(-\s*\d+(?:ST|ND|RD|TH)\s*FLOOR|\d+(?:ST|ND|RD|TH)\s*FLOOR)', clean, re.IGNORECASE)
    if floor_match:
        sheet_name = f"{floor_match.group(1).replace('-', '').strip()} Sched"
    else:
        sheet_name = clean.replace("SINGLE DUCT & REHEAT TERMINAL ", "").replace("SCEHDULE", "SCHED").replace("SCHEDULE", "SCHED").strip()

    sheet_name = sheet_name[:31] if sheet_name else f"Table_{index}"
    return sheet_name


def convert_tables_to_excel(
    file_path: str | Path,
    output_excel_path: str | Path | None = None,
    base_dir: Path | None = None,
    quiet: bool = False
) -> Path:
    """
    Extracts tabular data from a PDF or text file and writes each table to a separate tab in Excel in excel_files/.
    """
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError(
            "openpyxl is required for Excel export. Please install it:\n"
            "  • Via pip: pip install openpyxl --break-system-packages\n"
            "  • Via apt: sudo apt install python3-openpyxl"
        )

    script_dir = base_dir or Path(__file__).resolve().parent
    _, _, excel_dir = setup_directories(script_dir)

    input_file = Path(file_path).resolve()

    if output_excel_path is None:
        output_excel = excel_dir / f"{input_file.stem}.xlsx"
    else:
        output_excel = Path(output_excel_path).resolve()

    if not quiet:
        print(f"Extracting table data from '{input_file.name}'...")

    tables = extract_tables_from_pdf_or_text(input_file)

    if not tables:
        raise ValueError(f"No tables or schedule data found in '{input_file.name}'.")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    PRIMARY_COLOR = "1F497D"  # Navy blue
    ALT_ROW_COLOR = "F2F5F9"  # Light tint
    BORDER_COLOR = "D9D9D9"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    alt_fill = PatternFill(start_color=ALT_ROW_COLOR, end_color=ALT_ROW_COLOR, fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR)
    )

    used_sheet_names = set()

    for i, table in enumerate(tables, start=1):
        raw_title = table.get("title", f"Table_{i}")
        sheet_name = sanitize_sheet_title(raw_title, i)

        base_name = sheet_name
        counter = 1
        while sheet_name in used_sheet_names:
            suffix = f"_{counter}"
            sheet_name = f"{base_name[:31-len(suffix)]}{suffix}"
            counter += 1
        used_sheet_names.add(sheet_name)

        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        headers = table.get("headers", [])
        rows = table.get("rows", [])

        # Write header row
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Write data rows
        for row_idx, row_data in enumerate(rows, start=2):
            for col_idx, val in enumerate(row_data, start=1):
                clean_val = str(val).replace(",", "").strip() if val is not None else ""

                if clean_val.isdigit():
                    cell_val = int(clean_val)
                else:
                    try:
                        cell_val = float(clean_val)
                    except ValueError:
                        cell_val = val

                cell = ws.cell(row=row_idx, column=col_idx, value=cell_val)
                cell.border = thin_border

                if row_idx % 2 == 1:
                    cell.fill = alt_fill

                if isinstance(cell_val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(cell_val, float):
                        cell.number_format = "0.0"
                    else:
                        cell.number_format = "#,##0"
                elif str(val).strip() in ("X", "-"):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[row_idx].height = 20

        # Freeze top row
        ws.freeze_panes = "A2"

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value) if cell.value is not None else ""
                max_len = max(max_len, len(val_str))
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 35)

    output_excel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_excel))

    if not quiet:
        rel_excel = output_excel.relative_to(script_dir) if output_excel.is_relative_to(script_dir) else output_excel
        print(f"Excel Export Successful!")
        print(f"  • Output Excel: {rel_excel}")
        print(f"  • Total Tabs:   {len(tables)}")
        for t in tables:
            title_clean = sanitize_sheet_title(t['title'], 0)
            print(f"    - Tab: '{title_clean}' ({len(t['rows'])} rows)")

    return output_excel


def main():
    base_dir = Path(__file__).resolve().parent
    pdf_dir, text_dir, excel_dir = setup_directories(base_dir)

    parser = argparse.ArgumentParser(
        description="Convert PDF files to plain text and multi-tab Excel spreadsheets.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""Directory Structure:
  root/
  ├── pdf_to_excel.py
  ├── pdf_files/        (Input PDFs, e.g. page_1.pdf)
  ├── text_files/       (Output/Input text files, e.g. page_1.txt)
  └── excel_files/      (Output Excel files, e.g. page_1.xlsx)

Examples:
  # Process a specific file (e.g., page_1.pdf in pdf_files/):
  python3 pdf_to_excel.py page_1.pdf

  # Process all PDFs in pdf_files/:
  python3 pdf_to_excel.py

  # Only export Excel spreadsheet:
  python3 pdf_to_excel.py page_1.pdf --excel-only

  # Only export text file:
  python3 pdf_to_excel.py page_1.pdf --text-only
"""
    )

    parser.add_argument(
        "file_path",
        type=str,
        nargs="?",
        default=None,
        help="Path or name of the file to process (defaults to processing all PDFs in pdf_files/)."
    )
    parser.add_argument(
        "-o", "--output",
        type=str,
        default=None,
        help="Custom path for output text file."
    )
    parser.add_argument(
        "--excel-output",
        type=str,
        default=None,
        help="Custom path for output Excel file."
    )
    parser.add_argument(
        "--excel-only",
        action="store_true",
        help="Only generate the Excel file, skipping text output."
    )
    parser.add_argument(
        "--text-only",
        action="store_true",
        help="Only generate the text file, skipping Excel output."
    )
    parser.add_argument(
        "-e", "--engine",
        choices=["auto", "pypdf", "pdfplumber"],
        default="auto",
        help="PDF text extraction engine (default: auto)."
    )
    parser.add_argument(
        "-q", "--quiet",
        action="store_true",
        help="Suppress console status messages."
    )

    args = parser.parse_args()

    files_to_process = []

    if args.file_path:
        resolved_file = resolve_input_file(args.file_path, base_dir)
        files_to_process.append(resolved_file)
    else:
        pdf_candidates = list(pdf_dir.glob("*.pdf")) + list(pdf_dir.glob("*.PDF"))
        if pdf_candidates:
            files_to_process = sorted(pdf_candidates)
        else:
            text_candidates = list(text_dir.glob("*.txt"))
            if text_candidates:
                files_to_process = sorted(text_candidates)

    if not files_to_process:
        print(f"No PDF files found in '{pdf_dir}'. Please place PDF files in 'pdf_files/' and rerun.", file=sys.stderr)
        sys.exit(1)

    for current_file in files_to_process:
        try:
            if not args.excel_only and current_file.suffix.lower() == ".pdf":
                convert_pdf_to_text(
                    pdf_path=current_file,
                    output_path=args.output,
                    base_dir=base_dir,
                    engine=args.engine,
                    quiet=args.quiet
                )

            if not args.text_only:
                convert_tables_to_excel(
                    file_path=current_file,
                    output_excel_path=args.excel_output,
                    base_dir=base_dir,
                    quiet=args.quiet
                )

        except Exception as e:
            print(f"Error processing '{current_file.name}': {e}", file=sys.stderr)


if __name__ == "__main__":
    main()